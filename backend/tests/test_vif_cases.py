"""Casos de violencia intrafamiliar de las Comisarías: anonimización, catálogos y análisis."""
import io
from datetime import date

import pytest

from services import vif_cases as vif

SECRET = b"prueba"
TODAY = date(2026, 9, 26)


@pytest.fixture(autouse=True)
def no_geocoding(monkeypatch):
    known = {"BONANZA": "BONANZA", "TERRANOVA": "TERRANOVA"}
    monkeypatch.setattr(vif, "place", lambda value: (None, None, "SIN_DATO") if vif.blank(value) else (
        str(value), known.get(vif.norm(value), vif.norm(value)), "SI" if vif.norm(value) in known else "NO"))


def test_columnas_con_datos_personales_se_descartan():
    assert vif.column_kind("Nombre de la víctima") == "PERSONAL"
    assert vif.column_kind("No. Cédula") == "PERSONAL"
    assert vif.column_kind("Teléfono de contacto") == "PERSONAL"
    assert vif.column_kind("Dirección de residencia") == "PERSONAL"
    assert vif.column_kind("Fecha de nacimiento") == "PERSONAL"
    assert vif.column_kind("Observaciones") == "TEXTO_LIBRE"
    assert vif.column_kind("Relato de los hechos") == "TEXTO_LIBRE"
    assert vif.column_kind("Barrio") == "DATO"


def test_asignacion_automatica_de_columnas():
    columns = ["No. Radicado", "FECHA DE ATENCIÓN", "Barrio/Vereda", "Sexo", "Edad", "Tipo de Violencia",
               "Parentesco con el agresor", "Nivel de riesgo", "Medida de protección", "Fecha medida",
               "Seguimiento", "Reincidencia", "Otra cosa"]
    mapping = vif.auto_mapping(columns)
    assert mapping["codigo_caso"] == "No. Radicado"
    assert mapping["fecha_atencion"] == "FECHA DE ATENCIÓN"
    assert mapping["fecha_medida"] == "Fecha medida"
    assert mapping["lugar"] == "Barrio/Vereda"
    assert mapping["relacion"] == "Parentesco con el agresor"
    assert mapping["medida"] == "Medida de protección"
    assert "Otra cosa" not in mapping.values()
    # La asignación confirmada antes se respeta.
    assert vif.auto_mapping(columns, {"lugar": "Otra cosa"})["lugar"] == "Otra cosa"


def test_catalogos():
    assert vif.parse_date("05/03/2026") == date(2026, 3, 5)
    assert vif.parse_date("2026-03-05T00:00:00") == date(2026, 3, 5)
    assert vif.parse_date("46086") == date(2026, 3, 5)
    assert vif.parse_date("sin dato") is None
    assert [vif.age_group(v) for v in ("34", "18-28", "60 o más", "Adolescente", "adulto mayor", "", "x")] == [
        "ADULTEZ", "JUVENTUD", "PERSONA_MAYOR", "ADOLESCENCIA", "PERSONA_MAYOR", "SIN_DATO", "SIN_DATO"]
    assert vif.sex("F") == "MUJER" and vif.sex("Masculino") == "HOMBRE" and vif.sex("") == "SIN_DATO"
    # "M" depende del archivo: con F/M es masculino; con H/M es mujer; sola es ambigua.
    assert vif.m_meaning(["F", "M"]) == "HOMBRE" and vif.m_meaning(["H", "M"]) == "MUJER"
    assert vif.m_meaning(["M"]) is None and vif.sex("M") == "SIN_DATO"
    assert vif.sex("M", "HOMBRE") == "HOMBRE" and vif.sex("m", "MUJER") == "MUJER"
    assert vif.violence_types("Física y psicológica") == ["FISICA", "PSICOLOGICA"]
    assert vif.violence_types("verbal") == ["PSICOLOGICA"]
    assert vif.violence_types("maltrato") == ["OTRA"]
    assert [vif.relationship(v) for v in ("Compañero permanente", "Ex compañero", "Excónyuge", "Padrastro",
                                          "Hijo", "Tía", "Vecino", "")] == [
        "PAREJA", "EXPAREJA", "EXPAREJA", "PADRE_MADRE", "HIJO_HIJA", "OTRO_FAMILIAR", "OTRO", "SIN_DATO"]
    assert [vif.risk_level(v) for v in ("Riesgo extremo", "ALTO", "moderado", "variable", "Sin riesgo", "N/A")] == [
        "EXTREMO", "GRAVE", "MODERADO", "BAJO", "SIN_RIESGO", "NO_APLICA"]
    assert vif.measures("Desalojo y prohibición de acercarse") == ["DESALOJO", "ALEJAMIENTO"]
    assert vif.measures("Ninguna") == ["SIN_MEDIDA"]
    assert [vif.yes_no(v) for v in ("Sí", "X", "No", "0", "2", "12/03/2026", "No se realizó", "")] == [
        "SI", "SI", "NO", "NO", "SI", "SI", "NO", "SIN_DATO"]


def test_codigo_de_caso_es_irreversible_y_estable():
    key = vif.case_key("RAD-2026-001", SECRET)
    assert key and "2026" not in key and len(key) == 64
    assert vif.case_key(" rad-2026-001 ", SECRET) == key
    assert vif.case_key("RAD-2026-001", b"otra") != key
    assert vif.case_key("", SECRET) is None


MAPPING = {"codigo_caso": "Caso", "fecha_atencion": "Fecha", "lugar": "Barrio", "sexo": "Sexo", "edad": "Edad",
           "tipo_violencia": "Tipo", "relacion": "Relación", "antecedentes": "Antecedentes", "nivel_riesgo": "Riesgo",
           "medida": "Medida", "fecha_medida": "Fecha medida", "seguimiento": "Seguimiento", "reincidencia": "Reincidencia"}


def row(**values):
    base = {"Caso": "C1", "Fecha": "10/08/2026", "Barrio": "Bonanza", "Sexo": "Mujer", "Edad": "34", "Tipo": "Física",
            "Relación": "Esposo", "Antecedentes": "Sí", "Riesgo": "Alto", "Medida": "Desalojo",
            "Fecha medida": "11/08/2026", "Seguimiento": "Sí", "Reincidencia": "No"}
    base.update(values)
    return {key: value for key, value in base.items() if value is not None}


def test_procesar_filas_valida_y_resume():
    rows = [
        row(),
        row(),  # misma atención repetida en el archivo
        row(Caso="C2", Fecha="31/12/2030"),  # fecha futura
        row(Caso="C3", Fecha=None),
        row(Caso="C4", Barrio="Villa Sin Mapa", **{"Fecha medida": "01/08/2026"}),
    ]
    cases, summary = vif.process(rows, MAPPING, SECRET, TODAY)
    assert summary["total_rows"] == 5 and summary["accepted"] == 2
    assert summary["rejected"] == 2 and summary["duplicates"] == 1
    assert summary["with_warnings"] == 1
    assert summary["unknown_places"] == [{"lugar": "Villa Sin Mapa", "casos": 1}]
    assert [e["fila"] for e in summary["errors"]] == [4, 5]
    assert summary["has_case_code"] and summary["missing_fields"] == []
    first = cases[0]
    assert first["lugar"] == "BONANZA" and first["relacion"] == "PAREJA" and first["nivel_riesgo"] == "GRAVE"
    assert "C1" not in str(first.values())


def test_sin_codigo_la_identidad_es_el_contenido():
    mapping = {key: value for key, value in MAPPING.items() if key != "codigo_caso"}
    cases, summary = vif.process([row(), row(Tipo="Sexual")], mapping, SECRET, TODAY)
    assert len(cases) == 2 and summary["duplicates"] == 0 and not summary["has_case_code"]


def case(fecha, key=None, **values):
    base = {"entity": "Comisaría Primera de Familia", "case_key": key, "fecha_atencion": fecha, "lugar": "BONANZA",
            "lugar_reconocido": "SI", "sexo": "MUJER", "rango_edad": "ADULTEZ", "tipos_violencia": ["FISICA"],
            "relacion": "PAREJA", "antecedentes": "NO", "nivel_riesgo": "MODERADO", "medidas": ["DESALOJO"],
            "fecha_medida": fecha, "seguimiento": "SI", "reincidencia": "SIN_DATO"}
    base.update(values)
    return base


def test_analisis_reincidencia_factores_y_respuesta():
    cases = [
        case(date(2026, 1, 10), "A", antecedentes="SI"), case(date(2026, 3, 10), "A", antecedentes="SI"),
        case(date(2026, 2, 1), "B", antecedentes="SI", fecha_medida=date(2026, 2, 10)),
        case(date(2026, 2, 2), "C"), case(date(2026, 2, 3), "D"), case(date(2026, 2, 4), "E"),
        case(date(2026, 2, 5), "F"), case(date(2026, 2, 6), "G"),
        case(date(2026, 8, 1), None, lugar="TERRANOVA", reincidencia="SI", sexo="HOMBRE", seguimiento="NO",
             medidas=["SIN_MEDIDA"], fecha_medida=None),
    ]
    result = vif.analyze(cases, date(2026, 1, 1), date(2026, 8, 31), police={2026: 300})
    assert result["total"] == 9
    reinc = result["reincidence"]
    assert reinc["cases_by_code"] == 7 and reinc["repeated_by_code"] == 1 and reinc["median_days_to_repeat"] == 59
    assert reinc["flag_known"] == 1 and reinc["rate_by_flag"] == 100.0
    antecedentes = {row["code"]: row for row in result["factors"]["antecedentes"]}
    assert antecedentes["SI"]["cases"] == 3 and antecedentes["SI"]["small"] and antecedentes["SI"]["rate"] is None
    assert antecedentes["NO"]["cases"] == 6 and antecedentes["NO"]["repeats"] == 1
    response = result["response"]
    assert response["with_measure"] == 8 and response["median_days"] == 0 and response["over_7_days_pct"] == 12.5
    assert response["followup_pct"] == pytest.approx(88.9)
    assert result["territory"][0] == {"lugar": "BONANZA", "count": 8, "pct": 88.9, "recent": 0, "prior": 1,
                                      "recognized": True}
    assert result["police"] == [{"year": 2026, "comisarias": 9, "policia": 300}]
    assert result["women_pct"] == pytest.approx(88.9)
    assert [point["month"] for point in result["series"]] == ["2026-01", "2026-02", "2026-03", "2026-08"]


def test_analisis_sin_datos():
    assert vif.analyze([], date(2026, 1, 1), date(2026, 2, 1))["status"] == "SIN_DATOS"


def test_plantilla_excel():
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(vif.template_xlsx()))
    assert book.sheetnames == ["Casos", "Instrucciones"]
    headers = [cell.value for cell in book["Casos"][1]]
    assert headers[0] == "Código interno del caso" and "Tipo de violencia" in headers and len(headers) == 13
    # La plantilla se reconoce sola al cargarla.
    assert set(vif.auto_mapping(headers)) == set(vif.FIELDS)
