"""Seguimiento de compromisos de los Consejos de Seguridad.

- Importa la hoja "Seguimiento de Compromisos" (exportada a CSV o XLSX) sin pisar el estado
  que el equipo ya registró en el SISC: el acta es la fuente del compromiso; el SISC, la
  fuente de su avance.
- Calcula lo que el Consejo debe mirar: compromisos atrasados, repetidos sin cumplirse,
  sin fecha o sin información.
- Genera la "Lectura de compromisos" del orden del día.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from datetime import date, datetime
from typing import Any, Dict, Iterable, List, Optional

from sqlalchemy.orm import Session

from db.models_council import COMMITMENT_STATUSES, CouncilActRead, CouncilCommitment, CouncilCommitmentUpdate
from services.act_reader import INSTANCES, normalize, read_act
from services.council_topics import LABELS as TOPIC_LABELS
from services.council_topics import topics_in, theme_for

CLOSED_STATUSES = {"CUMPLIDO", "NO_CUMPLIDO", "DESCARTADO"}
STATUS_LABELS = {
    "SIN_INFORMACION": "Sin información",
    "PENDIENTE": "Pendiente",
    "EN_CURSO": "En curso",
    "CUMPLIDO": "Cumplido",
    "NO_CUMPLIDO": "No cumplido",
    "DESCARTADO": "Descartado",
}
REPEATED_THRESHOLD = 2

# Encabezados de la hoja -> campos del modelo (se comparan sin tildes ni mayúsculas).
SHEET_COLUMNS = {
    "ID": "code",
    "FECHA ORIGEN": "origin_date",
    "ACTA ORIGEN": "origin_act",
    "TIPO SESION": "session_type",
    "COMPROMISO CONSOLIDADO": "text",
    "RESPONSABLE PRINCIPAL": "responsible",
    "PLAZO TEXTUAL": "deadline_text",
    "FECHA LIMITE": "deadline_date",
    "ESTADO EN ACTA ORIGEN": "origin_status",
    "PRIORIDAD": "priority",
    "TEMA": "theme",
    "TERRITORIO": "territory",
    "MENCIONES": "mentions",
    "ULTIMA MENCION": "last_mention_date",
    "RELACIONADO CON": "related_to",
    "FUENTE ACTA": "source_act",
    "UBICACION FUENTE": "source_location",
    "NOTAS": "notes",
}
# Campos que vienen del acta y se actualizan en cada importación.
SOURCE_FIELDS = (
    "origin_date", "origin_act", "session_type", "text", "responsible", "deadline_text",
    "deadline_date", "theme", "territory", "mentions", "last_mention_date",
)


def _norm(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", text).strip().upper()


def parse_date(value: Any) -> Optional[date]:
    """Fechas de la hoja: dd/mm/aaaa (también aaaa-mm-dd o fecha de Excel)."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    match = re.search(r"\b(\d{1,2})/(\d{1,2})/(\d{4})\b", text)
    if match:
        day, month, year = (int(part) for part in match.groups())
        try:
            return date(year, month, day)
        except ValueError:
            return None
    match = re.search(r"\b(\d{4})-(\d{2})-(\d{2})\b", text)
    if match:
        try:
            return date(*(int(part) for part in match.groups()))
        except ValueError:
            return None
    return None


def _clean(value: Any) -> Optional[str]:
    text = str(value or "").strip()
    return text or None


def rows_from_table(rows: Iterable[List[Any]]) -> List[Dict[str, Any]]:
    """Convierte las filas de la hoja en compromisos (ubica sola la fila de encabezados)."""
    rows = [list(row) for row in rows]
    header_index = next(
        (i for i, row in enumerate(rows) if {"ID", "COMPROMISO CONSOLIDADO"} <= {_norm(cell) for cell in row}),
        None,
    )
    if header_index is None:
        raise ValueError("No se encontró la tabla de compromisos (columnas 'ID' y 'Compromiso consolidado').")
    headers = [SHEET_COLUMNS.get(_norm(cell)) for cell in rows[header_index]]
    commitments = []
    for row in rows[header_index + 1:]:
        record = {field: row[i] for i, field in enumerate(headers) if field and i < len(row)}
        code = _clean(record.get("code"))
        if not code or not _clean(record.get("text")):
            if commitments:
                break  # fin de la tabla
            continue
        mentions = re.sub(r"\D", "", str(record.get("mentions") or "1")) or "1"
        extra = {key: _clean(record.get(key)) for key in ("origin_status", "related_to", "source_act", "source_location")}
        commitments.append({
            "code": code,
            "origin_date": parse_date(record.get("origin_date")),
            "origin_act": _clean(record.get("origin_act")),
            "session_type": _clean(record.get("session_type")),
            "text": _clean(record.get("text")),
            "responsible": _clean(record.get("responsible")),
            "deadline_text": _clean(record.get("deadline_text")),
            "deadline_date": parse_date(record.get("deadline_date")),
            "priority": _clean(record.get("priority")),
            "theme": _clean(record.get("theme")),
            "territory": _clean(record.get("territory")),
            "mentions": int(mentions),
            "last_mention_date": parse_date(record.get("last_mention_date")),
            "source_ref": " · ".join(filter(None, [extra.get("source_act"), extra.get("source_location")])) or None,
            "notes": _clean(record.get("notes")),
            "extra": {key: value for key, value in extra.items() if value},
        })
    return commitments


def rows_from_file(content: bytes, filename: str) -> List[Dict[str, Any]]:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        return rows_from_table(csv.reader(io.StringIO(text)))
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            try:
                return rows_from_table(sheet.iter_rows(values_only=True))
            except ValueError:
                continue
        raise ValueError("Ninguna hoja del archivo tiene la tabla de compromisos.")
    raise ValueError("Formato no admitido. Exporte la hoja como CSV o XLSX.")


def import_commitments(db: Session, commitments: List[Dict[str, Any]], username: str) -> Dict[str, int]:
    """Crea los nuevos y actualiza los datos de origen de los existentes, sin tocar su estado."""
    created = updated = unchanged = 0
    for item in commitments:
        row = db.query(CouncilCommitment).filter(CouncilCommitment.code == item["code"]).first()
        if row is None:
            row = CouncilCommitment(**{key: value for key, value in item.items()}, status="SIN_INFORMACION",
                                    validated="PENDIENTE", version=1, updated_by=username)
            db.add(row)
            db.flush()
            db.add(CouncilCommitmentUpdate(commitment_id=row.id, version=1, action="IMPORTADO",
                                           new_status=row.status, note=item.get("origin_act"), username=username))
            created += 1
            continue
        changes = {field: item[field] for field in SOURCE_FIELDS if item.get(field) is not None and getattr(row, field) != item[field]}
        if not changes:
            unchanged += 1
            continue
        for field, value in changes.items():
            setattr(row, field, value)
        row.version += 1
        row.updated_by = username
        db.add(CouncilCommitmentUpdate(commitment_id=row.id, version=row.version, action="EDICION",
                                       previous_status=row.status, new_status=row.status,
                                       note="Actualizado desde la hoja de seguimiento: " + ", ".join(sorted(changes)),
                                       username=username))
        updated += 1
    db.commit()
    return {"created": created, "updated": updated, "unchanged": unchanged, "total": len(commitments)}


def update_status(db: Session, code: str, *, status: str, note: Optional[str], evidence_url: Optional[str],
                  deadline_date: Optional[date], expected_version: int, username: str) -> CouncilCommitment:
    if status not in COMMITMENT_STATUSES:
        raise ValueError("Estado no válido.")
    row = db.query(CouncilCommitment).filter(CouncilCommitment.code == code).with_for_update().first()
    if row is None:
        raise LookupError("Compromiso no encontrado.")
    if row.version != expected_version:
        raise PermissionError("Otra persona actualizó este compromiso. Recargue antes de guardar.")
    note = (note or "").strip() or None
    evidence_url = (evidence_url or "").strip() or None
    # Solo enlaces web: un "javascript:" o "data:" guardado aquí se ejecutaría al abrir el soporte.
    if evidence_url and not re.match(r"^https?://[^\s]+$", evidence_url, flags=re.IGNORECASE):
        raise ValueError("El enlace al soporte debe empezar por https:// o http://.")
    # Un compromiso cumplido o descartado debe decir cómo se sabe.
    if status in {"CUMPLIDO", "DESCARTADO", "NO_CUMPLIDO"} and not (note or evidence_url):
        raise ValueError("Para cerrar un compromiso escriba cómo se verificó (nota o enlace al soporte).")
    previous = row.status
    row.status = status
    if deadline_date is not None:
        row.deadline_date = deadline_date
    row.validated = "VALIDADO"
    row.version += 1
    row.updated_by = username
    db.add(CouncilCommitmentUpdate(commitment_id=row.id, version=row.version, action="ESTADO",
                                   previous_status=previous, new_status=status, note=note,
                                   evidence_url=evidence_url, username=username))
    db.commit()
    db.refresh(row)
    return row


def flags(row: CouncilCommitment, today: Optional[date] = None) -> List[str]:
    today = today or date.today()
    result = []
    open_ = row.status not in CLOSED_STATUSES
    if open_ and row.deadline_date and row.deadline_date < today:
        result.append("ATRASADO")
    if open_ and (row.mentions or 1) >= REPEATED_THRESHOLD:
        result.append("REPETIDO")
    if open_ and not row.deadline_date:
        result.append("SIN_FECHA")
    if row.status == "SIN_INFORMACION":
        result.append("SIN_INFORMACION")
    return result


def serialize(row: CouncilCommitment, today: Optional[date] = None) -> Dict[str, Any]:
    return {
        "code": row.code,
        "instance": row.instance or "CONSEJO_SEGURIDAD",
        "instance_label": INSTANCES.get(row.instance or "CONSEJO_SEGURIDAD", INSTANCES["OTRA"])["label"],
        "kind": row.kind or "COMPROMISO",
        "origin_date": row.origin_date.isoformat() if row.origin_date else None,
        "origin_act": row.origin_act,
        "session_type": row.session_type,
        "text": row.text,
        "responsible": row.responsible,
        "deadline_text": row.deadline_text,
        "deadline_date": row.deadline_date.isoformat() if row.deadline_date else None,
        "status": row.status,
        "status_label": STATUS_LABELS.get(row.status, row.status),
        "validated": row.validated,
        "theme": row.theme,
        "territory": row.territory,
        "mentions": row.mentions,
        "last_mention_date": row.last_mention_date.isoformat() if row.last_mention_date else None,
        "source_ref": row.source_ref,
        "version": row.version,
        "updated_by": row.updated_by,
        "updated_at": row.updated_at.isoformat() if row.updated_at else None,
        "flags": flags(row, today),
    }


def _priority_key(item: Dict[str, Any]):
    # Primero lo que el Consejo ha pedido más veces sin que se cumpla; luego los atrasados.
    return (-(item["mentions"] or 1), "ATRASADO" not in item["flags"], item["origin_date"] or "")


def summary(rows: List[CouncilCommitment], today: Optional[date] = None) -> Dict[str, Any]:
    items = [serialize(row, today) for row in rows]
    by_status = {status: 0 for status in COMMITMENT_STATUSES}
    for item in items:
        by_status[item["status"]] = by_status.get(item["status"], 0) + 1
    open_items = [item for item in items if item["status"] not in CLOSED_STATUSES]
    return {
        "total": len(items),
        "open": len(open_items),
        "by_status": by_status,
        "overdue": sum("ATRASADO" in item["flags"] for item in items),
        "repeated": sum("REPETIDO" in item["flags"] for item in items),
        "without_deadline": sum("SIN_FECHA" in item["flags"] for item in items),
        "without_information": by_status.get("SIN_INFORMACION", 0),
        "attention": sorted(
            [item for item in open_items if {"ATRASADO", "REPETIDO"} & set(item["flags"])], key=_priority_key
        )[:10],
    }


def agenda_text(rows: List[CouncilCommitment], today: Optional[date] = None) -> str:
    """Texto para el punto "Lectura de compromisos" del orden del día."""
    today = today or date.today()
    items = [serialize(row, today) for row in rows]
    open_items = sorted([item for item in items if item["status"] not in CLOSED_STATUSES], key=_priority_key)
    closed = [item for item in items if item["status"] == "CUMPLIDO"]
    lines = [f"LECTURA DE COMPROMISOS · corte {today.strftime('%d/%m/%Y')}",
             f"{len(items)} compromisos registrados: {len(open_items)} abiertos y {len(closed)} cumplidos.", ""]

    def line(item):
        details = [item["responsible"] or "Responsable por definir"]
        if item["deadline_date"]:
            details.append(f"plazo {datetime.fromisoformat(item['deadline_date']).strftime('%d/%m/%Y')}")
        elif item["deadline_text"]:
            details.append(f"plazo: {item['deadline_text']}")
        else:
            details.append("sin fecha límite")
        marks = []
        if "ATRASADO" in item["flags"]:
            marks.append("ATRASADO")
        if "REPETIDO" in item["flags"]:
            marks.append(f"pedido {item['mentions']} veces")
        mark = f" [{' · '.join(marks)}]" if marks else ""
        return f"- {item['code']}{mark}: {item['text']} ({'; '.join(details)}). Estado: {item['status_label']}."

    for title, selection in (
        ("1. Atrasados o repetidos sin cumplirse", [i for i in open_items if {"ATRASADO", "REPETIDO"} & set(i["flags"])]),
        ("2. Otros compromisos abiertos", [i for i in open_items if not ({"ATRASADO", "REPETIDO"} & set(i["flags"]))]),
        ("3. Cumplidos", closed),
    ):
        lines.append(title.upper())
        lines.extend(line(item) for item in selection) if selection else lines.append("- Ninguno.")
        lines.append("")
    without_deadline = sum("SIN_FECHA" in i["flags"] for i in open_items)
    if without_deadline:
        lines.append(f"Nota: {without_deadline} compromisos abiertos no tienen fecha límite. Se recomienda definirla en esta sesión.")
    return "\n".join(lines).strip()


def executive_block(rows: List[CouncilCommitment], today: Optional[date] = None) -> Dict[str, Any]:
    """Resumen para el Parte ejecutivo."""
    data = summary(rows, today)
    items = [serialize(row, today) for row in rows]
    results = [item for item in items if item["status"] == "CUMPLIDO"]
    # Los atrasados se listan aparte: en "attention" pueden quedar fuera detrás de los repetidos.
    overdue_items = sorted((item for item in items if "ATRASADO" in item["flags"]), key=_priority_key)
    return {**data, "results": results[:5], "overdue_items": overdue_items[:3]}


# ---------------------------------------------------------------------------
# Lectura de actas: propuestas para revisión humana y confirmación
# ---------------------------------------------------------------------------

_STOPWORDS = {
    "PARA", "CON", "LAS", "LOS", "DEL", "POR", "QUE", "UNA", "SUS", "SOBRE", "ENTRE", "ANTES", "COMO", "ESTE", "ESTA",
    "REALIZAR", "SECRETARIA", "SEGURIDAD", "CONVIVENCIA", "POLICIA", "NACIONAL", "MUNICIPIO", "JAMUNDI",
}


def _tokens(text: str) -> set:
    return {word for word in re.findall(r"[A-Z0-9]{4,}", normalize(text)) if word not in _STOPWORDS}


def similarity(a: str, b: str) -> float:
    left, right = _tokens(a), _tokens(b)
    if not left or not right:
        return 0.0
    return round(len(left & right) / len(left | right), 2)


def similar_commitments(text: str, rows: List[CouncilCommitment], limit: int = 3, threshold: float = 0.3) -> List[Dict[str, Any]]:
    scored = sorted(((similarity(text, row.text), row) for row in rows), key=lambda pair: -pair[0])
    return [
        {"code": row.code, "text": row.text, "instance": row.instance, "status": row.status, "score": score,
         "origin_date": row.origin_date.isoformat() if row.origin_date else None,
         "last_mention_date": row.last_mention_date.isoformat() if row.last_mention_date else None}
        for score, row in scored[:limit] if score >= threshold
    ]


SAME_SESSION_SCORE = 0.3


def _matching_text(text: str) -> str:
    """Quita el título corto antes de los dos puntos ("Podar árboles: Realizar la poda ...") para comparar."""
    head, sep, body = (text or "").partition(":")
    return body if sep and len(head.split()) <= 6 and len(body.strip()) > 15 else text


def _annotate_reading(db: Session, data: Dict[str, Any], sha: str) -> Dict[str, Any]:
    """Compara lo leído con lo registrado HOY: sirve al leer y al retomar una lectura pendiente."""
    existing = db.query(CouncilCommitment).all()
    act_date = data.get("act_date")
    for proposal in data["proposals"]:
        proposal["similar"] = similar_commitments(_matching_text(proposal["text"]), existing)
        top = proposal["similar"][0] if proposal["similar"] else None
        # Si ya existe uno muy parecido, se propone sumarle la mención en vez de duplicarlo.
        proposal["suggested_action"] = "LINK" if top and top["score"] >= 0.5 else "CREATE"
        # Si el parecido nació en esta misma acta (p. ej. ya vino de la hoja), sumarle mención lo contaría dos veces.
        # Dentro de la misma sesión basta un parecido menor: suele ser otra redacción del mismo compromiso.
        same_session = next((item for item in proposal["similar"]
                             if act_date and item["origin_date"] == act_date and item["score"] >= SAME_SESSION_SCORE), None)
        proposal["already_registered"] = bool(same_session)
        if same_session:
            proposal["similar"] = [same_session] + [item for item in proposal["similar"] if item is not same_session]
            proposal["suggested_action"] = "SKIP"
    for item in data["reread"]:
        similar = similar_commitments(item["text"], existing, limit=1, threshold=0.25)
        item["similar"] = similar
        item["already_counted"] = bool(similar and act_date and similar[0]["last_mention_date"]
                                       and similar[0]["last_mention_date"] >= act_date)
    data["warnings"] = [warning for warning in data["warnings"] if warning["code"] != "ACTA_YA_PROCESADA"]
    previous = db.query(CouncilActRead).filter(CouncilActRead.sha256 == sha, CouncilActRead.status == "CONFIRMADA").first()
    if previous:
        data["warnings"].insert(0, {
            "code": "ACTA_YA_PROCESADA",
            "message": f"Esta misma acta ya se confirmó el {previous.confirmed_at.strftime('%d/%m/%Y') if previous.confirmed_at else 'antes'}; confirmarla de nuevo sumaría menciones repetidas.",
        })
    return data


def _read_payload(record: CouncilActRead, data: Dict[str, Any]) -> Dict[str, Any]:
    return {"read_id": str(record.id), "filename": record.filename, "created_by": record.created_by,
            "created_at": record.created_at.isoformat() if record.created_at else None, **data}


def read_act_for_review(db: Session, content: bytes, filename: str, instance_hint: Optional[str], username: str,
                        use_ocr: bool = False) -> Dict[str, Any]:
    import hashlib

    sha = hashlib.sha256(content).hexdigest()
    # La misma acta subida dos veces sin confirmar: se retoma la lectura pendiente en vez de duplicarla.
    pending = db.query(CouncilActRead).filter(CouncilActRead.sha256 == sha, CouncilActRead.status == "PENDIENTE").first()
    if pending and (not instance_hint or instance_hint == pending.instance):
        return open_act_read(db, str(pending.id))
    reading, reread = read_act(content, filename, instance_hint, use_ocr=use_ocr)
    data = reading.to_dict()
    data["proposals"] = [{**item.__dict__, "index": index} for index, item in enumerate(reading.proposals)]
    data["reread"] = [{"index": index, "text": text} for index, text in enumerate(reread)]
    data = _annotate_reading(db, data, sha)
    record = CouncilActRead(
        filename=filename[:255], sha256=sha, instance=reading.instance, act_number=reading.act_number,
        act_date=date.fromisoformat(reading.act_date) if reading.act_date else None,
        reading=data, status="PENDIENTE", created_by=username,
    )
    db.add(record)
    db.commit()
    db.refresh(record)
    return _read_payload(record, data)


def _read_key(read_id: str):
    from uuid import UUID

    try:
        return UUID(str(read_id))
    except ValueError as error:
        raise LookupError("Lectura de acta no encontrada.") from error


def _get_read(db: Session, read_id: str) -> CouncilActRead:
    record = db.query(CouncilActRead).filter(CouncilActRead.id == _read_key(read_id)).first()
    if record is None:
        raise LookupError("Lectura de acta no encontrada.")
    return record


def open_act_read(db: Session, read_id: str) -> Dict[str, Any]:
    record = _get_read(db, read_id)
    data = dict(record.reading)
    if record.status == "PENDIENTE":
        data = _annotate_reading(db, data, record.sha256)
    return {**_read_payload(record, data), "status": record.status, "result": record.result}


def list_act_reads(db: Session, status: Optional[str] = "PENDIENTE") -> List[Dict[str, Any]]:
    query = db.query(CouncilActRead)
    if status:
        query = query.filter(CouncilActRead.status == status)
    rows = query.order_by(CouncilActRead.act_date.asc().nullslast(), CouncilActRead.created_at.asc()).all()
    return [{
        "read_id": str(row.id), "filename": row.filename, "instance": row.instance,
        "instance_label": INSTANCES.get(row.instance, INSTANCES["OTRA"])["label"],
        "act_number": row.act_number, "act_date": row.act_date.isoformat() if row.act_date else None,
        "status": row.status, "proposals": len((row.reading or {}).get("proposals", [])),
        "warnings": [warning["code"] for warning in (row.reading or {}).get("warnings", [])],
        "created_by": row.created_by, "created_at": row.created_at.isoformat() if row.created_at else None,
    } for row in rows]


def store_historical_act(db: Session, content: bytes, filename: str, username: str,
                         use_ocr: bool = False) -> Dict[str, Any]:
    """Guarda un acta de años anteriores solo para análisis: no pasa por la bandeja ni crea compromisos."""
    import hashlib

    sha = hashlib.sha256(content).hexdigest()
    existing = db.query(CouncilActRead).filter(CouncilActRead.sha256 == sha).first()
    if existing:
        return {"read_id": str(existing.id), "status": existing.status, "duplicate": True, "filename": existing.filename,
                "act_number": existing.act_number, "act_date": existing.act_date.isoformat() if existing.act_date else None}
    reading, reread = read_act(content, filename, use_ocr=use_ocr)
    data = reading.to_dict()
    data["proposals"] = [{**item.__dict__, "index": index} for index, item in enumerate(reading.proposals)]
    data["reread"] = [{"index": index, "text": text} for index, text in enumerate(reread)]
    record = CouncilActRead(
        filename=filename[:255], sha256=sha, instance=reading.instance, act_number=reading.act_number,
        act_date=date.fromisoformat(reading.act_date) if reading.act_date else None,
        reading=data, status="HISTORICA", created_by=username,
    )
    db.add(record)
    db.commit()
    db.refresh(record)
    return {"read_id": str(record.id), "status": record.status, "duplicate": False, "filename": record.filename,
            "instance": record.instance, "act_number": record.act_number,
            "act_date": reading.act_date, "proposals": len(data["proposals"]),
            "topics": sorted(data.get("topics", {}), key=lambda key: -data["topics"][key])[:5],
            "warnings": [warning["code"] for warning in data["warnings"]]}


# Un asunto cuenta como "tratado" en una sesión si aparece en al menos este número de párrafos
# (una mención de paso no es un tema de la sesión) o si hay un compromiso sobre él.
DISCUSSED_MIN_PARAGRAPHS = 2
CLOSED_STATUSES = {"CUMPLIDO", "NO_CUMPLIDO", "DESCARTADO"}


def recurrence(db: Session, reads: Optional[List[CouncilActRead]] = None,
               commitments: Optional[List[CouncilCommitment]] = None) -> Dict[str, Any]:
    """Asuntos que vuelven sesión tras sesión y año tras año, con lo que sigue abierto hoy."""
    reads = reads if reads is not None else db.query(CouncilActRead).all()
    commitments = commitments if commitments is not None else db.query(CouncilCommitment).all()
    sessions: Dict[tuple, Dict[str, Any]] = {}

    def session(instance: str, when: Optional[date], fallback: str) -> Dict[str, Any]:
        # Misma instancia y fecha = misma sesión, aunque haya dos archivos o venga también de la hoja.
        key = (instance, when.isoformat() if when else fallback)
        return sessions.setdefault(key, {"instance": instance, "date": when, "discussed": set(), "committed": {},
                                         "sources": set()})

    for read in reads:
        data = read.reading or {}
        item = session(read.instance, read.act_date, str(read.id))
        item["sources"].add(read.status)
        item["discussed"].update(key for key, count in (data.get("topics") or {}).items() if count >= DISCUSSED_MIN_PARAGRAPHS)
        for proposal in data.get("proposals", []):
            for key in topics_in(proposal.get("text", "")):
                item["committed"].setdefault(key, proposal.get("text"))
    for row in commitments:
        item = session(row.instance or "CONSEJO_SEGURIDAD", row.origin_date, row.code)
        item["sources"].add("HOJA")
        for key in topics_in(row.text):
            item["committed"].setdefault(key, row.text)

    today_open: Dict[str, List[CouncilCommitment]] = {}
    for row in commitments:
        if row.status not in CLOSED_STATUSES:
            for key in topics_in(row.text):
                today_open.setdefault(key, []).append(row)

    topics: Dict[str, Dict[str, Any]] = {}
    for item in sessions.values():
        for key in item["discussed"] | set(item["committed"]):
            topic = topics.setdefault(key, {"key": key, "label": TOPIC_LABELS[key], "dates": [], "committed_dates": [],
                                            "instances": set(), "examples": []})
            if item["date"]:
                topic["dates"].append(item["date"])
            topic["instances"].add(item["instance"])
            if key in item["committed"]:
                if item["date"]:
                    topic["committed_dates"].append(item["date"])
                topic["examples"].append((item["date"], item["instance"], item["committed"][key]))

    rows = []
    for topic in topics.values():
        dates = sorted(topic["dates"])
        years: Dict[str, int] = {}
        for when in dates:
            years[str(when.year)] = years.get(str(when.year), 0) + 1
        examples = sorted((example for example in topic["examples"] if example[0]), key=lambda example: example[0])
        chosen = examples[:1] + examples[-1:] if len(examples) > 1 else examples
        open_rows = sorted(today_open.get(topic["key"], []), key=lambda row: -(row.mentions or 1))
        rows.append({
            "key": topic["key"], "label": topic["label"],
            "sessions": len(dates), "sessions_with_commitment": len(topic["committed_dates"]),
            "years": years, "first_date": dates[0].isoformat() if dates else None,
            "last_date": dates[-1].isoformat() if dates else None,
            "instances": sorted(INSTANCES.get(code, INSTANCES["OTRA"])["label"] for code in topic["instances"]),
            "open_commitments": len(open_rows), "open_codes": [row.code for row in open_rows[:4]],
            "examples": [{"date": when.isoformat(), "instance": INSTANCES.get(inst, INSTANCES["OTRA"])["label"], "text": text[:280]}
                         for when, inst, text in chosen],
        })
    rows = [row for row in rows if row["sessions"] >= 3 or len(row["years"]) >= 2]
    rows.sort(key=lambda row: (-len(row["years"]), -row["sessions"], row["label"]))

    coverage: Dict[str, Dict[str, int]] = {}
    for item in sessions.values():
        if item["date"]:
            label = INSTANCES.get(item["instance"], INSTANCES["OTRA"])["label"]
            coverage.setdefault(str(item["date"].year), {}).setdefault(label, 0)
            coverage[str(item["date"].year)][label] += 1
    unreadable = [read.filename for read in reads
                  if any(warning.get("code") == "SIN_TEXTO" for warning in (read.reading or {}).get("warnings", []))]
    return {"topics": rows, "coverage": dict(sorted(coverage.items())), "sessions": len(sessions),
            "historical_reads": sum(1 for read in reads if read.status == "HISTORICA"), "unreadable": unreadable}


def discard_act_read(db: Session, read_id: str) -> None:
    record = _get_read(db, read_id)
    if record.status != "PENDIENTE":
        raise PermissionError("Solo se puede descartar una lectura sin confirmar.")
    db.delete(record)
    db.commit()


def _next_code(db: Session, instance: str, year: int) -> str:
    prefix = INSTANCES.get(instance, INSTANCES["OTRA"])["prefix"]
    stem = f"{prefix}-{year}-"
    codes = [row[0] for row in db.query(CouncilCommitment.code).filter(CouncilCommitment.code.like(f"{stem}%")).all()]
    numbers = [int(code[len(stem):]) for code in codes if code[len(stem):].isdigit()]
    return f"{stem}{(max(numbers) if numbers else 0) + 1:03d}"


def confirm_act(db: Session, read_id: str, decisions: List[Dict[str, Any]], reread_links: List[Dict[str, Any]],
                username: str) -> Dict[str, Any]:
    """Aplica lo que la persona decidió: crear, sumar mención a uno existente o descartar."""
    record = db.query(CouncilActRead).filter(CouncilActRead.id == _read_key(read_id)).with_for_update().first()
    if record is None:
        raise LookupError("Lectura de acta no encontrada.")
    if record.status == "CONFIRMADA":
        raise PermissionError("Esta lectura ya fue confirmada.")
    if record.status != "PENDIENTE":
        raise PermissionError("Las actas históricas son solo para análisis; no generan compromisos.")
    proposals = {item["index"]: item for item in record.reading.get("proposals", [])}
    act_date = record.act_date
    act_label = f"Acta {record.act_number}" if record.act_number else record.filename
    instance_label = INSTANCES.get(record.instance, INSTANCES["OTRA"])["label"]
    created, linked, skipped = [], [], 0

    def link(code: str, note: str) -> bool:
        row = db.query(CouncilCommitment).filter(CouncilCommitment.code == code).first()
        if row is None:
            return False
        if code in linked:  # una misma acta cuenta como una sola mención
            return True
        row.mentions = (row.mentions or 1) + 1
        if act_date and (not row.last_mention_date or act_date > row.last_mention_date):
            row.last_mention_date = act_date
        row.version += 1
        row.updated_by = username
        db.add(CouncilCommitmentUpdate(commitment_id=row.id, version=row.version, action="EDICION",
                                       previous_status=row.status, new_status=row.status, note=note, username=username))
        linked.append(code)
        return True

    for decision in decisions:
        proposal = proposals.get(decision.get("index"))
        if proposal is None:
            raise ValueError(f"La propuesta {decision.get('index')} no pertenece a esta acta.")
        action = decision.get("action")
        if action == "SKIP":
            skipped += 1
            continue
        if action == "LINK":
            if not link(decision.get("link_code") or "", f"Mencionado de nuevo en {act_label} ({instance_label})."):
                raise ValueError(f"No existe el compromiso {decision.get('link_code')}.")
            continue
        if action != "CREATE":
            raise ValueError("Acción no válida.")
        text = (decision.get("text") or proposal["text"] or "").strip()
        if not text:
            raise ValueError("El compromiso no puede quedar vacío.")
        deadline = decision.get("deadline_date") or proposal.get("deadline_date")
        code = _next_code(db, record.instance, (act_date or date.today()).year)
        row = CouncilCommitment(
            code=code, instance=record.instance, kind=proposal.get("kind") or "COMPROMISO",
            origin_date=act_date, origin_act=act_label, session_type=instance_label, text=text,
            responsible=(decision.get("responsible") if decision.get("responsible") is not None else proposal.get("responsible")) or None,
            deadline_text=proposal.get("deadline_text"),
            deadline_date=date.fromisoformat(deadline) if deadline else None,
            status="SIN_INFORMACION", validated="VALIDADO", mentions=1, last_mention_date=act_date, theme=theme_for(text),
            source_ref=f"{record.filename} · {proposal.get('method')}",
            extra={"read_id": str(record.id), "flags": proposal.get("flags", []), "source": proposal.get("source")},
            version=1, updated_by=username,
        )
        db.add(row)
        db.flush()
        db.add(CouncilCommitmentUpdate(commitment_id=row.id, version=1, action="IMPORTADO", new_status=row.status,
                                       note=f"Leído de {act_label} ({instance_label}) y confirmado.", username=username))
        created.append(code)

    for item in reread_links:
        if item.get("link_code"):
            link(item["link_code"], f"Releído en la lectura de compromisos de {act_label} ({instance_label}).")

    record.status = "CONFIRMADA"
    record.confirmed_by = username
    record.confirmed_at = datetime.utcnow()
    record.result = {"created": created, "linked": linked, "skipped": skipped}
    db.commit()
    return record.result
