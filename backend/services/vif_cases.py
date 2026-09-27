"""Casos de violencia intrafamiliar de las Comisarías de Familia: lectura, anonimización y análisis.

Cada comisaría entrega su base en su propio formato. El SISC reconoce las columnas por su
nombre (la persona confirma o corrige la asignación), descarta las columnas con datos
personales y normaliza los valores a catálogos comunes. El análisis responde lo que pidió el
Observatorio: tendencia, concentración territorial, reincidencia, factores de riesgo y
oportunidad de la respuesta institucional. Uso interno y reservado.
"""
from __future__ import annotations

import hashlib
import hmac
import io
import math
import os
import re
import secrets
import statistics
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from typing import Any, Dict, Iterable, List, Optional, Tuple

from sqlalchemy.orm import Session

from services.file_reader import normalize_sheet_token

MIN_YEAR = 2010
MIN_CELL = 5  # por debajo, una tasa no se interpreta (pocos casos)
MAX_ERRORS_SHOWN = 20

# campo: (etiqueta, obligatorio, alias de encabezado)
FIELDS: Dict[str, Tuple[str, bool, List[str]]] = {
    "codigo_caso": ("Código interno del caso (opcional, se guarda como huella)", False,
                    ["CODIGO CASO", "CODIGO DEL CASO", "CODIGO", "RADICADO", "NO RADICADO", "NUMERO RADICADO",
                     "HISTORIA", "NO HISTORIA", "NUMERO DE HISTORIA", "NUMERO CASO", "NO CASO", "CASO", "ID CASO",
                     "EXPEDIENTE", "NO EXPEDIENTE", "NUMERO EXPEDIENTE", "CONSECUTIVO CASO"]),
    "fecha_atencion": ("Fecha de atención", True,
                       ["FECHA DE ATENCION", "FECHA ATENCION", "FECHA", "FECHA RECEPCION", "FECHA DE RECEPCION",
                        "FECHA INGRESO", "FECHA DE INGRESO", "FECHA DENUNCIA", "FECHA SOLICITUD", "FECHA RADICACION"]),
    "lugar": ("Barrio, sector o corregimiento", False,
              ["BARRIO SECTOR O CORREGIMIENTO", "BARRIO", "BARRIO SECTOR", "SECTOR", "CORREGIMIENTO", "VEREDA",
               "BARRIO VEREDA", "LUGAR", "UBICACION", "LOCALIDAD", "BARRIO O VEREDA", "ZONA BARRIO"]),
    "sexo": ("Sexo de la víctima", False, ["SEXO", "SEXO VICTIMA", "SEXO DE LA VICTIMA", "GENERO", "GENERO VICTIMA"]),
    "edad": ("Edad o rango de edad de la víctima", False,
             ["RANGO DE EDAD", "RANGO EDAD", "EDAD", "EDAD VICTIMA", "EDAD DE LA VICTIMA", "GRUPO ETARIO",
              "CICLO VITAL", "RANGO ETARIO", "RANGO DE EDAD DE LA VICTIMA"]),
    "tipo_violencia": ("Tipo de violencia", False,
                       ["TIPO DE VIOLENCIA", "TIPO VIOLENCIA", "CLASE DE VIOLENCIA", "VIOLENCIA", "TIPOLOGIA",
                        "MODALIDAD", "FORMA DE VIOLENCIA"]),
    "relacion": ("Relación entre víctima y presunto agresor", False,
                 ["RELACION VICTIMA AGRESOR", "RELACION ENTRE VICTIMA Y PRESUNTO AGRESOR", "RELACION CON EL AGRESOR",
                  "RELACION", "PARENTESCO", "VINCULO", "PARENTESCO CON EL AGRESOR", "RELACION AGRESOR"]),
    "antecedentes": ("Antecedentes o episodios previos", False,
                     ["ANTECEDENTES O EPISODIOS PREVIOS", "ANTECEDENTES", "EPISODIOS PREVIOS", "HECHOS PREVIOS",
                      "DENUNCIAS PREVIAS"]),
    "nivel_riesgo": ("Nivel de riesgo", False,
                     ["NIVEL DE RIESGO", "NIVEL RIESGO", "RIESGO", "VALORACION DE RIESGO", "VALORACION RIESGO"]),
    "medida": ("Medida de protección adoptada", False,
               ["MEDIDA DE PROTECCION ADOPTADA", "MEDIDA DE PROTECCION", "MEDIDA PROTECCION", "MEDIDA ADOPTADA",
                "MEDIDAS DE PROTECCION", "MEDIDA", "MEDIDAS"]),
    "fecha_medida": ("Fecha de la medida", False,
                     ["FECHA DE LA MEDIDA", "FECHA MEDIDA", "FECHA MEDIDA DE PROTECCION", "FECHA DE MEDIDA",
                      "FECHA MEDIDA PROTECCION"]),
    "seguimiento": ("Seguimiento realizado", False, ["SEGUIMIENTO REALIZADO", "SEGUIMIENTO", "SEGUIMIENTOS"]),
    "reincidencia": ("Nuevos episodios o reincidencia", False,
                     ["REGISTRO DE NUEVOS EPISODIOS O REINCIDENCIA", "REINCIDENCIA", "NUEVOS EPISODIOS",
                      "NUEVO EPISODIO", "REINCIDENTE", "NUEVOS HECHOS"]),
}
RECOMMENDED = ["lugar", "sexo", "edad", "tipo_violencia", "relacion", "nivel_riesgo", "medida"]

# Encabezados que delatan datos personales: la columna se descarta al leer.
PII_TOKENS = ["NOMBRE", "APELLIDO", "CEDULA", "DOCUMENTO", "IDENTIFICACION", "NUMEROID", "NOID", "NUIP",
              "TARJETADEIDENTIDAD", "PASAPORTE", "TELEFONO", "CELULAR", "CORREO", "EMAIL", "DIRECCION", "FIRMA",
              "WHATSAPP", "NACIMIENTO"]
# Texto libre: puede traer nombres o relatos; tampoco se guarda.
FREE_TEXT_TOKENS = ["OBSERVACION", "DESCRIPCION", "RELATO", "NARRACION", "COMENTARIO", "DETALLE", "HECHOSNARRADOS"]

AGE_GROUPS = [("PRIMERA_INFANCIA", 0, 5, "0 a 5 años"), ("INFANCIA", 6, 11, "6 a 11 años"),
              ("ADOLESCENCIA", 12, 17, "12 a 17 años"), ("JUVENTUD", 18, 28, "18 a 28 años"),
              ("ADULTEZ", 29, 59, "29 a 59 años"), ("PERSONA_MAYOR", 60, 130, "60 años o más")]

LABELS = {
    "sexo": {"MUJER": "Mujer", "HOMBRE": "Hombre", "INTERSEXUAL": "Intersexual", "SIN_DATO": "Sin dato"},
    "rango_edad": {**{code: label for code, _, _, label in AGE_GROUPS}, "SIN_DATO": "Sin dato"},
    "tipos_violencia": {"FISICA": "Física", "PSICOLOGICA": "Psicológica", "SEXUAL": "Sexual",
                        "ECONOMICA": "Económica", "PATRIMONIAL": "Patrimonial", "NEGLIGENCIA": "Negligencia o abandono",
                        "OTRA": "Otra"},
    "relacion": {"PAREJA": "Pareja", "EXPAREJA": "Expareja", "PADRE_MADRE": "Padre, madre o padrastros",
                 "HIJO_HIJA": "Hijo o hija", "HERMANO": "Hermano o hermana", "OTRO_FAMILIAR": "Otro familiar",
                 "OTRO": "Otra relación", "SIN_DATO": "Sin dato"},
    "nivel_riesgo": {"SIN_RIESGO": "Sin riesgo", "BAJO": "Bajo o variable", "MODERADO": "Moderado",
                     "GRAVE": "Grave o alto", "EXTREMO": "Extremo", "NO_APLICA": "No aplica", "SIN_DATO": "Sin dato"},
    "medidas": {"DESALOJO": "Desalojo del agresor", "ALEJAMIENTO": "Alejamiento o prohibición de acercarse",
                "TRATAMIENTO": "Tratamiento terapéutico o reeducativo", "PROTECCION_POLICIVA": "Protección policiva",
                "CUSTODIA_ALIMENTOS": "Custodia, alimentos o visitas", "CONMINACION": "Conminación o amonestación",
                "OTRA": "Otra medida", "SIN_MEDIDA": "Sin medida"},
    "si_no": {"SI": "Sí", "NO": "No", "SIN_DATO": "Sin dato"},
}


# ---------------------------------------------------------------- normalización de valores

def norm(value: Any) -> str:
    text = unicodedata.normalize("NFD", str(value if value is not None else "").upper())
    text = "".join(c for c in text if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", text).strip()


def blank(value: Any) -> bool:
    return norm(value) in {"", "NAN", "NONE", "NULL", "NAT", "-", "N/D", "ND", "SIN DATO", "SIN INFORMACION",
                           "NO REPORTA", "NO REGISTRA"}


def parse_date(value: Any) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if blank(value):
        return None
    text = str(value).strip()
    match = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", text)
    if match:
        year, month, day = map(int, match.groups())
    else:
        match = re.match(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})$", text.split(" ")[0])
        if match:
            day, month, year = map(int, match.groups())
            year += 2000 if year < 100 else 0
        else:
            try:
                serial = float(text)
            except ValueError:
                return None
            if 20000 <= serial <= 80000:  # número de serie de Excel
                return date(1899, 12, 30) + timedelta(days=int(serial))
            return None
    try:
        return date(year, month, day)
    except ValueError:
        return None


def age_group(value: Any) -> str:
    if blank(value):
        return "SIN_DATO"
    text = norm(value)
    for pattern, code in (("PRIMERA INFANCIA", "PRIMERA_INFANCIA"), ("ADULTO MAYOR", "PERSONA_MAYOR"),
                          ("PERSONA MAYOR", "PERSONA_MAYOR"), ("INFANCIA", "INFANCIA"), ("NINEZ", "INFANCIA"),
                          ("ADOLESC", "ADOLESCENCIA"), ("JUVENTUD", "JUVENTUD"), ("JOVEN", "JUVENTUD"),
                          ("ADULT", "ADULTEZ")):
        if pattern in text:
            return code
    number = re.search(r"\d+(?:[.,]\d+)?", text)
    if not number:
        return "SIN_DATO"
    years = int(float(number.group().replace(",", ".")))  # un rango se ubica por su límite inferior
    for code, low, high, _ in AGE_GROUPS:
        if low <= years <= high:
            return code
    return "SIN_DATO"


def sex(value: Any, m_means: Optional[str] = None) -> str:
    """m_means: qué significa "M" en este archivo (MUJER si usa H/M, HOMBRE si usa F/M)."""
    text = norm(value)
    if text == "M":
        return m_means or "SIN_DATO"
    if text in {"F", "MUJER", "FEMENINO", "FEMENINA"} or text.startswith(("MUJER", "FEMEN")):
        return "MUJER"
    if text in {"H", "HOMBRE", "MASCULINO", "MASCULINA"} or text.startswith(("HOMBRE", "MASCUL")):
        return "HOMBRE"
    if text.startswith("INTERSEX"):
        return "INTERSEXUAL"
    return "SIN_DATO"


def m_meaning(values: Iterable[Any]) -> Optional[str]:
    """Deduce del archivo qué significa la letra M en la columna de sexo."""
    letters = {norm(value) for value in values}
    if letters & {"F", "FEMENINO", "FEMENINA"}:
        return "HOMBRE"
    if letters & {"H", "HOMBRE"}:
        return "MUJER"
    return None


def violence_types(value: Any) -> List[str]:
    if blank(value):
        return []
    text = norm(value)
    found = []
    for code, patterns in (("FISICA", ["FISIC"]), ("PSICOLOGICA", ["PSICOL", "VERBAL", "EMOCIONAL"]),
                           ("SEXUAL", ["SEXUAL"]), ("ECONOMICA", ["ECONOM"]), ("PATRIMONIAL", ["PATRIMON"]),
                           ("NEGLIGENCIA", ["NEGLIG", "ABANDONO", "DESCUIDO"])):
        if any(pattern in text for pattern in patterns):
            found.append(code)
    return found or ["OTRA"]


def relationship(value: Any) -> str:
    if blank(value):
        return "SIN_DATO"
    text = norm(value)
    if re.search(r"\bEX\b|\bEX-|EXPAREJA|EXCOMPA|EXESPOS|EXNOVI|EXCONYUG|EXMARIDO", text):
        return "EXPAREJA"
    rules = (
        ("PAREJA", ["ESPOS", "COMPANER", "CONYUG", "NOVI", "PAREJA", "MARIDO", "SENTIMENTAL"]),
        ("PADRE_MADRE", ["PADRASTRO", "MADRASTRA", "PADRE", "MADRE", "PAPA", "MAMA", "PROGENITOR"]),
        ("HIJO_HIJA", ["HIJASTR", "HIJO", "HIJA"]),
        ("HERMANO", ["HERMAN"]),
        ("OTRO_FAMILIAR", ["TIO", "TIA", "PRIMO", "PRIMA", "ABUEL", "NIET", "SUEGR", "CUNAD", "YERNO", "NUERA",
                           "SOBRIN", "FAMILIAR"]),
    )
    for code, patterns in rules:
        if any(re.search(r"\b" + pattern, text) for pattern in patterns):
            return code
    return "OTRO"


def risk_level(value: Any) -> str:
    if blank(value):
        return "SIN_DATO"
    text = norm(value)
    if "SIN RIESGO" in text or text == "NINGUNO":
        return "SIN_RIESGO"
    for code, patterns in (("EXTREMO", ["EXTREM"]), ("GRAVE", ["GRAVE", "ALTO", "SEVER"]),
                           ("MODERADO", ["MODERAD", "MEDIO"]), ("BAJO", ["BAJO", "LEVE", "VARIABLE"]),
                           ("NO_APLICA", ["NO APLICA", "N/A", "NA"])):
        if any(re.search(r"\b" + re.escape(pattern), text) for pattern in patterns):
            return code
    return "SIN_DATO"


def measures(value: Any) -> List[str]:
    if blank(value):
        return []
    text = norm(value)
    if text in {"NO", "NINGUNA", "NINGUNO", "SIN MEDIDA", "NO SE ADOPTO", "NO APLICA"} or text.startswith("SIN MEDIDA"):
        return ["SIN_MEDIDA"]
    found = []
    for code, patterns in (("DESALOJO", ["DESALOJ"]),
                           ("ALEJAMIENTO", ["ABSTEN", "ALEJ", "ACERCA", "PROHIB", "INGRES"]),
                           ("TRATAMIENTO", ["TRATAMIENTO", "TERAP", "REEDUCA", "PSICOLOG"]),
                           ("PROTECCION_POLICIVA", ["POLIC", "PROTECCION TEMPORAL", "ACOMPANAMIENTO"]),
                           ("CUSTODIA_ALIMENTOS", ["CUSTODIA", "CUIDADO PERSONAL", "ALIMENT", "VISITA"]),
                           ("CONMINACION", ["CONMIN", "AMONEST"])):
        if any(pattern in text for pattern in patterns):
            found.append(code)
    return found or ["OTRA"]


def yes_no(value: Any) -> str:
    if blank(value):
        return "SIN_DATO"
    text = norm(value)
    if text in {"NO", "N", "0", "FALSE", "FALSO", "NINGUNO", "NINGUNA", "NO TIENE", "SIN ANTECEDENTES"} \
            or text.startswith("NO "):
        return "NO"
    try:
        return "SI" if float(text.replace(",", ".")) > 0 else "NO"
    except ValueError:
        return "SI"  # "Sí", "X", una fecha o una anotación: hubo algo que registrar


# ---------------------------------------------------------------- columnas del archivo

def column_kind(header: str) -> str:
    """PERSONAL (se descarta), TEXTO_LIBRE (se descarta) o DATO."""
    token = normalize_sheet_token(header)
    if any(pii in token for pii in PII_TOKENS):
        return "PERSONAL"
    if any(free in token for free in FREE_TEXT_TOKENS):
        return "TEXTO_LIBRE"
    return "DATO"


def auto_mapping(columns: List[str], previous: Optional[Dict[str, str]] = None) -> Dict[str, str]:
    """Asigna columnas a campos: primero la asignación confirmada antes, luego coincidencias exactas
    y luego parciales (el alias más largo gana). Cada columna y cada campo se usan una sola vez."""
    mapping: Dict[str, str] = {}
    used = set()
    for field, column in (previous or {}).items():
        if field in FIELDS and column in columns and column not in used:
            mapping[field] = column
            used.add(column)
    candidates = []
    for column in columns:
        token = normalize_sheet_token(column)
        if not token:
            continue
        for field, (_, _, aliases) in FIELDS.items():
            for alias in aliases:
                alias_token = normalize_sheet_token(alias)
                if token == alias_token:
                    candidates.append((2, len(alias_token), field, column))
                elif len(alias_token) >= 5 and alias_token in token:
                    candidates.append((1, len(alias_token), field, column))
    for _, _, field, column in sorted(candidates, reverse=True):
        if field not in mapping and column not in used:
            mapping[field] = column
            used.add(column)
    return mapping


def read_table(content: bytes, filename: str) -> Tuple[str, List[str], List[Dict[str, Any]]]:
    """Lee el libro y devuelve (hoja, columnas, filas como dict columna→texto)."""
    from services.file_reader import select_sheet_frame

    alias_map = {field: aliases for field, (_, _, aliases) in FIELDS.items()}
    sheet, frame = select_sheet_frame(content, filename, alias_map)
    frame = frame.dropna(how="all")
    columns = [str(column) for column in frame.columns]
    rows = []
    for record in frame.to_dict(orient="records"):
        row = {}
        for column, value in record.items():
            if value is None or (isinstance(value, float) and math.isnan(value)):
                continue
            if isinstance(value, datetime):
                value = value.date().isoformat()
            elif isinstance(value, date):
                value = value.isoformat()
            elif isinstance(value, float) and value.is_integer():
                value = str(int(value))
            else:
                value = str(value).strip()
            if value != "":
                row[str(column)] = value
        if row:
            rows.append(row)
    return str(sheet), columns, rows


# ---------------------------------------------------------------- filas

def case_key_secret(db: Optional[Session]) -> bytes:
    configured = os.getenv("VIF_CASE_KEY_SECRET", "").strip()
    if configured:
        return configured.encode()
    from db.models_vif import VifSetting

    row = db.get(VifSetting, "case_key_salt")
    if row is None:
        row = VifSetting(key="case_key_salt", value=secrets.token_hex(32))
        db.add(row)
        db.flush()
    return row.value.encode()


def case_key(code: Any, secret: bytes) -> Optional[str]:
    if blank(code):
        return None
    return hmac.new(secret, norm(code).encode(), hashlib.sha256).hexdigest()


def place(value: Any) -> Tuple[Optional[str], Optional[str], str]:
    """(original, nombre para análisis, reconocido SI|NO|SIN_DATO)."""
    from services.geocoding_service import GeocodingService

    if blank(value):
        return None, None, "SIN_DATO"
    original = str(value).strip()[:150]
    territory = GeocodingService.get_official_territory(original)
    if territory and territory.get("name"):
        return original, territory["name"], "SI"
    return original, norm(original)[:150], "NO"


def normalize_row(raw: Dict[str, Any], mapping: Dict[str, str], secret: bytes,
                  today: date, m_means: Optional[str] = None) -> Tuple[Optional[Dict[str, Any]], List[str], List[str]]:
    """Devuelve (caso, errores, advertencias). Con errores, la fila no se carga."""
    value = lambda field: raw.get(mapping[field]) if field in mapping else None
    errors, warnings = [], []
    attended = parse_date(value("fecha_atencion"))
    if attended is None:
        errors.append("sin fecha de atención válida")
    elif attended.year < MIN_YEAR or attended > today:
        errors.append(f"fecha de atención imposible ({attended:%d/%m/%Y})")
    measure_date = parse_date(value("fecha_medida"))
    if measure_date and (measure_date > today or measure_date.year < MIN_YEAR):
        warnings.append("fecha de la medida imposible: se omite")
        measure_date = None
    if measure_date and attended and measure_date < attended:
        warnings.append("la medida es anterior a la atención")
    if errors:
        return None, errors, warnings
    original, lugar, recognized = place(value("lugar"))
    if recognized == "NO":
        warnings.append("barrio o sector no reconocido")
    if norm(value("sexo")) == "M" and not m_means:
        warnings.append("sexo «M» ambiguo: no se sabe si es mujer o masculino")
    case = {
        "case_key": case_key(value("codigo_caso"), secret),
        "fecha_atencion": attended,
        "lugar_original": original,
        "lugar": lugar,
        "lugar_reconocido": recognized,
        "sexo": sex(value("sexo"), m_means),
        "rango_edad": age_group(value("edad")),
        "tipos_violencia": violence_types(value("tipo_violencia")),
        "relacion": relationship(value("relacion")),
        "antecedentes": yes_no(value("antecedentes")),
        "nivel_riesgo": risk_level(value("nivel_riesgo")),
        "medidas": measures(value("medida")),
        "fecha_medida": measure_date,
        "seguimiento": yes_no(value("seguimiento")),
        "reincidencia": yes_no(value("reincidencia")),
    }
    case["record_key"] = record_key(case)
    return case, errors, warnings


def record_key(case: Dict[str, Any]) -> str:
    """Con código de caso: una atención por caso y fecha. Sin código: la huella del contenido."""
    if case["case_key"]:
        raw = f"{case['case_key']}|{case['fecha_atencion']}"
    else:
        raw = "|".join(str(case[key]) for key in ("fecha_atencion", "lugar_original", "sexo", "rango_edad",
                                                   "tipos_violencia", "relacion", "antecedentes", "nivel_riesgo",
                                                   "medidas", "fecha_medida", "seguimiento", "reincidencia"))
    return hashlib.sha256(raw.encode()).hexdigest()


def process(rows: List[Dict[str, Any]], mapping: Dict[str, str], secret: bytes,
            today: Optional[date] = None) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    today = today or date.today()
    cases: Dict[str, Dict[str, Any]] = {}
    errors, unknown_places = [], Counter()
    rejected = duplicates = with_warnings = 0
    m_means = m_meaning(raw.get(mapping["sexo"]) for raw in rows) if "sexo" in mapping else None
    for index, raw in enumerate(rows, start=2):  # fila 1 = encabezado
        case, row_errors, warnings = normalize_row(raw, mapping, secret, today, m_means)
        if row_errors:
            rejected += 1
            if len(errors) < MAX_ERRORS_SHOWN:
                errors.append({"fila": index, "motivo": "; ".join(row_errors)})
            continue
        if warnings:
            with_warnings += 1
        if case["lugar_reconocido"] == "NO":
            unknown_places[case["lugar_original"]] += 1
        if case["record_key"] in cases:
            duplicates += 1
        cases[case["record_key"]] = case
    accepted = list(cases.values())
    dates = [case["fecha_atencion"] for case in accepted]
    summary = {
        "total_rows": len(rows), "accepted": len(accepted), "rejected": rejected, "duplicates": duplicates,
        "with_warnings": with_warnings, "errors": errors,
        "unknown_places": [{"lugar": name, "casos": count} for name, count in unknown_places.most_common(15)],
        "missing_fields": [field for field in RECOMMENDED if field not in mapping],
        "has_case_code": "codigo_caso" in mapping,
        "period_start": min(dates).isoformat() if dates else None,
        "period_end": max(dates).isoformat() if dates else None,
    }
    return accepted, summary


# ---------------------------------------------------------------- entregas

def _serialize_delivery(row, include_columns: bool = False) -> Dict[str, Any]:
    data = {
        "id": str(row.id), "entity": row.entity, "filename": row.filename, "sheet": row.sheet,
        "status": row.status, "mapping": row.mapping, "dropped_columns": row.dropped_columns,
        "ignored_columns": row.ignored_columns, "summary": row.summary,
        "period_start": row.period_start.isoformat() if row.period_start else None,
        "period_end": row.period_end.isoformat() if row.period_end else None,
        "created_by": row.created_by, "created_at": row.created_at.isoformat() if row.created_at else None,
        "decided_by": row.decided_by, "decided_at": row.decided_at.isoformat() if row.decided_at else None,
    }
    if include_columns and row.staged_rows is not None:
        columns = sorted({column for raw in row.staged_rows for column in raw})
        data["columns"] = columns
        hidden = row.mapping.get("codigo_caso")  # el código del caso no se muestra, ni de ejemplo
        data["samples"] = {column: sorted({raw[column] for raw in row.staged_rows[:200] if column in raw})[:4]
                           for column in columns if column != hidden}
    return data


def _last_mapping(db: Session, entity: str) -> Optional[Dict[str, str]]:
    from db.models_vif import VifDelivery

    row = (db.query(VifDelivery).filter(VifDelivery.entity == entity, VifDelivery.status == "CONFIRMADA")
           .order_by(VifDelivery.decided_at.desc()).first())
    return row.mapping if row else None


def preview(db: Session, content: bytes, filename: str, entity: str, username: str) -> Dict[str, Any]:
    from db.models_vif import ENTITIES, VifDelivery

    if entity not in ENTITIES:
        raise ValueError("Comisaría desconocida.")
    sheet, columns, rows = read_table(content, filename)
    if not rows:
        raise ValueError("El archivo no tiene filas con datos.")
    kinds = {column: column_kind(column) for column in columns}
    dropped = [column for column in columns if kinds[column] != "DATO"]
    allowed = [column for column in columns if kinds[column] == "DATO"]
    staged = [{column: value for column, value in raw.items() if kinds.get(column) == "DATO"} for raw in rows]
    mapping = auto_mapping(allowed, _last_mapping(db, entity))
    secret = case_key_secret(db)
    _, summary = process(staged, mapping, secret)
    sha = hashlib.sha256(content).hexdigest()
    earlier = db.query(VifDelivery).filter(VifDelivery.sha256 == sha, VifDelivery.status == "CONFIRMADA").first()
    summary["already_loaded"] = earlier.created_at.isoformat() if earlier else None
    delivery = VifDelivery(
        entity=entity, filename=filename[:255], sha256=sha, sheet=sheet[:120], mapping=mapping,
        dropped_columns=[{"column": column, "reason": kinds[column]} for column in dropped],
        ignored_columns=[column for column in allowed if column not in mapping.values()],
        staged_rows=staged, summary=summary, status="PREVISUALIZADA", created_by=username,
        period_start=date.fromisoformat(summary["period_start"]) if summary["period_start"] else None,
        period_end=date.fromisoformat(summary["period_end"]) if summary["period_end"] else None,
    )
    db.add(delivery)
    db.commit()
    db.refresh(delivery)
    return _serialize_delivery(delivery, include_columns=True)


def _pending(db: Session, delivery_id: str):
    from db.models_vif import VifDelivery
    import uuid

    try:
        row = db.get(VifDelivery, uuid.UUID(str(delivery_id)))
    except ValueError:
        row = None
    if row is None:
        raise LookupError("Entrega no encontrada.")
    if row.status != "PREVISUALIZADA":
        raise PermissionError(f"La entrega ya está {row.status.lower()}.")
    return row


def remap(db: Session, delivery_id: str, mapping: Dict[str, Optional[str]]) -> Dict[str, Any]:
    row = _pending(db, delivery_id)
    columns = {column for raw in row.staged_rows for column in raw}
    clean = {}
    for field, column in mapping.items():
        if field not in FIELDS:
            raise ValueError(f"Campo desconocido: {field}")
        if column:
            if column not in columns:
                raise ValueError(f"La columna «{column}» no está en el archivo o contiene datos personales.")
            if column in clean.values():
                raise ValueError(f"La columna «{column}» está asignada a dos campos.")
            clean[field] = column
    _, summary = process(row.staged_rows, clean, case_key_secret(db))
    summary["already_loaded"] = row.summary.get("already_loaded")
    row.mapping, row.summary = clean, summary
    row.ignored_columns = sorted(column for column in columns if column not in clean.values())
    row.period_start = date.fromisoformat(summary["period_start"]) if summary["period_start"] else None
    row.period_end = date.fromisoformat(summary["period_end"]) if summary["period_end"] else None
    db.commit()
    db.refresh(row)
    return _serialize_delivery(row, include_columns=True)


def confirm(db: Session, delivery_id: str, username: str) -> Dict[str, Any]:
    from db.models_vif import VifCase

    row = _pending(db, delivery_id)
    if "fecha_atencion" not in row.mapping:
        raise ValueError("Asigne la columna de la fecha de atención antes de confirmar.")
    cases, summary = process(row.staged_rows, row.mapping, case_key_secret(db))
    if not cases:
        raise ValueError("Ninguna fila es válida: revise la asignación de columnas.")
    existing = {case.record_key: case for case in db.query(VifCase).filter(
        VifCase.entity == row.entity, VifCase.record_key.in_([case["record_key"] for case in cases])).all()}
    created = updated = 0
    for case in cases:
        target = existing.get(case["record_key"])
        if target is None:
            target = VifCase(entity=row.entity, record_key=case["record_key"])
            db.add(target)
            created += 1
        else:
            updated += 1
        for key, value in case.items():
            if key != "record_key":
                setattr(target, key, value)
        target.delivery_id = row.id
    summary.update({"created": created, "updated": updated, "already_loaded": row.summary.get("already_loaded")})
    row.summary, row.status, row.staged_rows = summary, "CONFIRMADA", None
    row.decided_by, row.decided_at = username, datetime.utcnow()
    db.commit()
    db.refresh(row)
    return _serialize_delivery(row)


def discard(db: Session, delivery_id: str, username: str) -> Dict[str, Any]:
    row = _pending(db, delivery_id)
    row.status, row.staged_rows = "DESCARTADA", None
    row.decided_by, row.decided_at = username, datetime.utcnow()
    db.commit()
    db.refresh(row)
    return _serialize_delivery(row)


def list_deliveries(db: Session, limit: int = 50) -> List[Dict[str, Any]]:
    from db.models_vif import VifDelivery

    rows = db.query(VifDelivery).order_by(VifDelivery.created_at.desc()).limit(limit).all()
    return [_serialize_delivery(row) for row in rows]


def open_delivery(db: Session, delivery_id: str) -> Dict[str, Any]:
    from db.models_vif import VifDelivery
    import uuid

    try:
        row = db.get(VifDelivery, uuid.UUID(str(delivery_id)))
    except ValueError:
        row = None
    if row is None:
        raise LookupError("Entrega no encontrada.")
    return _serialize_delivery(row, include_columns=True)


# ---------------------------------------------------------------- análisis

def _rate(numerator: int, denominator: int) -> Optional[float]:
    return round(numerator / denominator * 100, 1) if denominator else None


def _counts(values: Iterable[str], labels: Dict[str, str]) -> List[Dict[str, Any]]:
    counter = Counter(values)
    total = sum(counter.values())
    return [{"code": code, "label": labels.get(code, code), "count": count, "pct": _rate(count, total)}
            for code, count in counter.most_common()]


def _month(value: date) -> str:
    return f"{value.year}-{value.month:02d}"


def reincidence(cases: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Reincidencia por código de caso (misma huella con más de una atención) y por la columna de la comisaría."""
    by_key = defaultdict(list)
    for case in cases:
        if case["case_key"]:
            by_key[case["case_key"]].append(case["fecha_atencion"])
    repeated = {key: sorted(dates) for key, dates in by_key.items() if len(set(dates)) > 1}
    gaps = [(dates[1] - dates[0]).days for dates in repeated.values()]
    flagged = [case for case in cases if case["reincidencia"] != "SIN_DATO"]
    return {
        "with_code": sum(1 for case in cases if case["case_key"]),
        "cases_by_code": len(by_key),
        "repeated_by_code": len(repeated),
        "rate_by_code": _rate(len(repeated), len(by_key)),
        "median_days_to_repeat": statistics.median(gaps) if gaps else None,
        "flag_known": len(flagged),
        "flag_yes": sum(1 for case in flagged if case["reincidencia"] == "SI"),
        "rate_by_flag": _rate(sum(1 for case in flagged if case["reincidencia"] == "SI"), len(flagged)),
    }


def is_repeat(case: Dict[str, Any], repeated_keys: set) -> Optional[bool]:
    if case["case_key"] in repeated_keys or case["reincidencia"] == "SI":
        return True
    if case["reincidencia"] == "NO" or case["case_key"]:
        return False
    return None


def factor_table(cases: List[Dict[str, Any]], field: str, labels: Dict[str, str], repeated_keys: set) -> List[Dict[str, Any]]:
    groups = defaultdict(lambda: [0, 0])
    for case in cases:
        repeat = is_repeat(case, repeated_keys)
        if repeat is None or case[field] == "SIN_DATO":
            continue
        groups[case[field]][0] += 1
        groups[case[field]][1] += int(repeat)
    return [{"code": code, "label": labels.get(code, code), "cases": n, "repeats": r,
             "rate": _rate(r, n) if n >= MIN_CELL else None, "small": n < MIN_CELL}
            for code, (n, r) in sorted(groups.items(), key=lambda item: -item[1][0])]


def analyze(cases: List[Dict[str, Any]], start: date, end: date, police: Optional[Dict[int, int]] = None) -> Dict[str, Any]:
    """Todo el análisis sobre casos ya normalizados (dicts). La reincidencia por código mira toda la historia."""
    repeated_keys = set()
    history = defaultdict(set)
    for case in cases:
        if case["case_key"] and case["fecha_atencion"] <= end:
            history[case["case_key"]].add(case["fecha_atencion"])
    repeated_keys = {key for key, dates in history.items() if len(dates) > 1}
    window = [case for case in cases if start <= case["fecha_atencion"] <= end]
    if not window:
        return {"status": "SIN_DATOS", "start": start.isoformat(), "end": end.isoformat(), "total": 0}

    months = defaultdict(Counter)
    for case in window:
        months[_month(case["fecha_atencion"])][case["entity"]] += 1
    series = [{"month": month, "total": sum(counter.values()), "by_entity": dict(counter)}
              for month, counter in sorted(months.items())]

    # Concentración territorial y cambio: últimos 90 días frente a los 90 anteriores.
    recent_start, prior_start = end - timedelta(days=89), end - timedelta(days=179)
    places = Counter(case["lugar"] for case in window if case["lugar"])
    recent = Counter(case["lugar"] for case in cases if case["lugar"] and recent_start <= case["fecha_atencion"] <= end)
    prior = Counter(case["lugar"] for case in cases
                    if case["lugar"] and prior_start <= case["fecha_atencion"] < recent_start)
    located = sum(places.values())
    territory = [{"lugar": name, "count": count, "pct": _rate(count, located), "recent": recent[name],
                  "prior": prior[name], "recognized": any(c["lugar"] == name and c["lugar_reconocido"] == "SI"
                                                          for c in window)}
                 for name, count in places.most_common(15)]

    both = [(case["fecha_medida"] - case["fecha_atencion"]).days for case in window
            if case["fecha_medida"] and case["fecha_medida"] >= case["fecha_atencion"]]
    with_measure = [case for case in window if case["medidas"] and case["medidas"] != ["SIN_MEDIDA"]]
    follow_known = [case for case in window if case["seguimiento"] != "SIN_DATO"]

    completeness = {}
    for field, empty in (("lugar", None), ("sexo", "SIN_DATO"), ("rango_edad", "SIN_DATO"), ("tipos_violencia", []),
                         ("relacion", "SIN_DATO"), ("antecedentes", "SIN_DATO"), ("nivel_riesgo", "SIN_DATO"),
                         ("medidas", []), ("fecha_medida", None), ("seguimiento", "SIN_DATO"),
                         ("reincidencia", "SIN_DATO")):
        completeness[field] = _rate(sum(1 for case in window if case[field] not in (empty, None)), len(window))

    by_year = Counter(case["fecha_atencion"].year for case in window)
    police_cross = [{"year": year, "comisarias": by_year[year], "policia": (police or {}).get(year)}
                    for year in sorted(by_year)]

    women = sum(1 for case in window if case["sexo"] == "MUJER")
    sex_age = defaultdict(Counter)
    for case in window:
        sex_age[case["rango_edad"]][case["sexo"]] += 1
    return {
        "status": "OK", "start": start.isoformat(), "end": end.isoformat(), "total": len(window),
        "by_entity": dict(Counter(case["entity"] for case in window)),
        "women_pct": _rate(women, sum(1 for case in window if case["sexo"] != "SIN_DATO")),
        "series": series,
        "territory": territory, "territory_located": located, "territory_unlocated": len(window) - located,
        "recent_window": {"start": recent_start.isoformat(), "prior_start": prior_start.isoformat()},
        "violence": _counts((code for case in window for code in case["tipos_violencia"]), LABELS["tipos_violencia"]),
        "relationship": _counts((case["relacion"] for case in window), LABELS["relacion"]),
        "sex": _counts((case["sexo"] for case in window), LABELS["sexo"]),
        "age": [{"code": code, "label": label, "count": sum(sex_age[code].values()),
                 "mujer": sex_age[code]["MUJER"], "hombre": sex_age[code]["HOMBRE"]}
                for code, _, _, label in AGE_GROUPS] +
               [{"code": "SIN_DATO", "label": "Sin dato", "count": sum(sex_age["SIN_DATO"].values()),
                 "mujer": sex_age["SIN_DATO"]["MUJER"], "hombre": sex_age["SIN_DATO"]["HOMBRE"]}],
        "risk": _counts((case["nivel_riesgo"] for case in window), LABELS["nivel_riesgo"]),
        "measures": _counts((code for case in window for code in case["medidas"]), LABELS["medidas"]),
        "response": {
            "with_measure": len(with_measure), "with_measure_pct": _rate(len(with_measure), len(window)),
            "with_both_dates": len(both),
            "median_days": statistics.median(both) if both else None,
            "within_1_day_pct": _rate(sum(1 for days in both if days <= 1), len(both)),
            "over_7_days_pct": _rate(sum(1 for days in both if days > 7), len(both)),
            "followup_known": len(follow_known),
            "followup_pct": _rate(sum(1 for case in follow_known if case["seguimiento"] == "SI"), len(follow_known)),
        },
        "reincidence": reincidence(window),
        "factors": {
            "antecedentes": factor_table(window, "antecedentes", LABELS["si_no"], repeated_keys),
            "nivel_riesgo": factor_table(window, "nivel_riesgo", LABELS["nivel_riesgo"], repeated_keys),
            "relacion": factor_table(window, "relacion", LABELS["relacion"], repeated_keys),
        },
        "completeness": completeness,
        "police": police_cross,
        "min_cell": MIN_CELL,
    }


def _case_dict(row) -> Dict[str, Any]:
    return {key: getattr(row, key) for key in ("entity", "case_key", "fecha_atencion", "lugar", "lugar_reconocido",
                                               "sexo", "rango_edad", "tipos_violencia", "relacion", "antecedentes",
                                               "nivel_riesgo", "medidas", "fecha_medida", "seguimiento",
                                               "reincidencia")}


def police_vif_by_year(db: Session) -> Dict[int, int]:
    """Denuncias de violencia intrafamiliar en Jamundí según Mindefensa (solo total municipal anual)."""
    from sqlalchemy import text

    rows = db.execute(text("""
        SELECT EXTRACT(YEAR FROM fecha_hecho)::int AS year, SUM(cantidad)::int AS total
        FROM national_crime_stats
        WHERE source_id = 'MINDEFENSA_MUNICIPAL_TOTAL' AND tipo_delito ILIKE '%intrafamiliar%'
          AND UPPER(municipio_normalizado) LIKE 'JAMUND%'
        GROUP BY 1
    """)).fetchall()
    return {row.year: row.total for row in rows if row.year}


def analysis(db: Session, entity: Optional[str] = None, start: Optional[date] = None,
             end: Optional[date] = None) -> Dict[str, Any]:
    from db.models_vif import VifCase

    query = db.query(VifCase)
    if entity:
        query = query.filter(VifCase.entity == entity)
    cases = [_case_dict(row) for row in query.all()]
    if not cases:
        return {"status": "SIN_DATOS", "total": 0, "deliveries": list_deliveries(db, 10)}
    last = max(case["fecha_atencion"] for case in cases)
    end = end or last
    start = start or (end - timedelta(days=364))
    result = analyze(cases, start, end, police_vif_by_year(db))
    result.update({"entity": entity, "data_first": min(c["fecha_atencion"] for c in cases).isoformat(),
                   "data_last": last.isoformat(), "deliveries": list_deliveries(db, 10)})
    return result


# ---------------------------------------------------------------- plantilla

TEMPLATE_VALUES = {
    "sexo": "Mujer · Hombre · Intersexual",
    "edad": "Edad en años (p. ej. 34) o rango: 0-5, 6-11, 12-17, 18-28, 29-59, 60 o más",
    "tipo_violencia": "Física · Psicológica · Sexual · Económica · Patrimonial · Negligencia. Varias separadas por coma",
    "relacion": "Pareja · Expareja · Padre/Madre · Hijo(a) · Hermano(a) · Otro familiar · Otra",
    "antecedentes": "Sí · No",
    "nivel_riesgo": "Sin riesgo · Bajo · Moderado · Grave · Extremo · No aplica",
    "medida": "Desalojo · Alejamiento · Tratamiento · Protección policiva · Custodia/alimentos · Conminación · Ninguna",
    "seguimiento": "Sí · No (o la fecha del seguimiento)",
    "reincidencia": "Sí · No (o el número de nuevos episodios)",
    "fecha_atencion": "dd/mm/aaaa",
    "fecha_medida": "dd/mm/aaaa",
    "lugar": "Nombre del barrio, sector, vereda o corregimiento (sin dirección exacta)",
    "codigo_caso": "Número interno de caso o historia. No escriba cédula ni nombre. El SISC lo convierte en un código irreversible",
}
TEMPLATE_HEADERS = {
    "codigo_caso": "Código interno del caso", "fecha_atencion": "Fecha de atención",
    "lugar": "Barrio, sector o corregimiento", "sexo": "Sexo de la víctima", "edad": "Rango de edad de la víctima",
    "tipo_violencia": "Tipo de violencia", "relacion": "Relación víctima - presunto agresor",
    "antecedentes": "Antecedentes o episodios previos", "nivel_riesgo": "Nivel de riesgo",
    "medida": "Medida de protección adoptada", "fecha_medida": "Fecha de la medida",
    "seguimiento": "Seguimiento realizado", "reincidencia": "Nuevos episodios o reincidencia",
}
TEMPLATE_LISTS = {
    "sexo": ["Mujer", "Hombre", "Intersexual"],
    "edad": ["0-5", "6-11", "12-17", "18-28", "29-59", "60 o más"],
    "relacion": ["Pareja", "Expareja", "Padre/Madre", "Hijo(a)", "Hermano(a)", "Otro familiar", "Otra"],
    "antecedentes": ["Sí", "No"], "seguimiento": ["Sí", "No"], "reincidencia": ["Sí", "No"],
    "nivel_riesgo": ["Sin riesgo", "Bajo", "Moderado", "Grave", "Extremo", "No aplica"],
}


def template_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    book = Workbook()
    sheet = book.active
    sheet.title = "Casos"
    fields = list(TEMPLATE_HEADERS)
    header_fill = PatternFill("solid", fgColor="281FD0")
    for index, field in enumerate(fields, start=1):
        cell = sheet.cell(row=1, column=index, value=TEMPLATE_HEADERS[field])
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = 24
        if field in TEMPLATE_LISTS:
            validation = DataValidation(type="list", formula1='"' + ",".join(TEMPLATE_LISTS[field]) + '"',
                                        allow_blank=True, showErrorMessage=False)
            sheet.add_data_validation(validation)
            validation.add(f"{get_column_letter(index)}2:{get_column_letter(index)}5000")
        if field.startswith("fecha"):
            for row in range(2, 5001):
                sheet.cell(row=row, column=index).number_format = "DD/MM/YYYY"
    sheet.row_dimensions[1].height = 42
    sheet.freeze_panes = "A2"

    guide = book.create_sheet("Instrucciones")
    guide.column_dimensions["A"].width = 38
    guide.column_dimensions["B"].width = 90
    rows = [
        ("Observatorio del Delito · Secretaría de Seguridad y Convivencia de Jamundí", ""),
        ("Una fila por atención. Si ya tienen su propia base, pueden enviarla tal como está: el SISC la adapta.", ""),
        ("NO incluya nombres, cédulas, teléfonos, correos ni direcciones exactas.", ""),
        ("", ""),
        ("Columna", "Qué escribir"),
    ] + [(TEMPLATE_HEADERS[field], TEMPLATE_VALUES[field]) for field in fields]
    for row in rows:
        guide.append(row)
    for cell in (guide["A1"], guide["A3"], guide["A5"], guide["B5"]):
        cell.font = Font(bold=True)
    for row in guide.iter_rows(min_row=6):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def catalog() -> Dict[str, Any]:
    from db.models_vif import ENTITIES

    return {"entities": list(ENTITIES),
            "fields": [{"code": field, "label": label, "required": required} for field, (label, required, _) in FIELDS.items()],
            "labels": LABELS, "min_cell": MIN_CELL}
